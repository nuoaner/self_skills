#!/usr/bin/env python3
import argparse
import json
from pathlib import Path
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

EMU = 914400.0


def text_of(shape):
    if getattr(shape, "has_text_frame", False):
        return (shape.text or "").strip()
    return ""


def classify(slide, stats):
    title = stats["title"].lower()
    if stats["index"] == 1:
        return "cover"
    if stats["chart_count"]:
        return "chart"
    if stats["table_count"]:
        return "table"
    if stats["arrow_count"] >= 2:
        return "process_or_diagram"
    if stats["picture_count"] >= 1 and stats["text_chars"] < 180:
        return "visual_explanation"
    if len(slide.shapes) <= 4 and stats["text_chars"] < 120:
        return "section_or_transition"
    if any(k in title for k in ("summary", "recap", "conclusion")):
        return "summary"
    return "content"


def inventory(path):
    prs = Presentation(path)
    slides = []
    for idx, slide in enumerate(prs.slides, 1):
        texts = [text_of(s) for s in slide.shapes if text_of(s)]
        title = ""
        if slide.shapes.title is not None:
            title = text_of(slide.shapes.title)
        if not title and texts:
            title = texts[0].splitlines()[0][:120]
        picture_count = sum(1 for s in slide.shapes if s.shape_type == MSO_SHAPE_TYPE.PICTURE)
        chart_count = sum(1 for s in slide.shapes if getattr(s, "has_chart", False))
        table_count = sum(1 for s in slide.shapes if getattr(s, "has_table", False))
        arrow_count = 0
        for s in slide.shapes:
            try:
                xml = s._element.xml
                if "headEnd" in xml or "tailEnd" in xml:
                    arrow_count += 1
            except Exception:
                pass
        try:
            notes = slide.notes_slide.notes_text_frame.text.strip()
        except Exception:
            notes = ""
        item = {
            "index": idx,
            "layout": getattr(slide.slide_layout, "name", "") or "",
            "title": title,
            "shape_count": len(slide.shapes),
            "picture_count": picture_count,
            "chart_count": chart_count,
            "table_count": table_count,
            "arrow_count": arrow_count,
            "text_chars": sum(len(t) for t in texts),
            "has_notes": bool(notes),
        }
        item["archetype"] = classify(slide, item)
        slides.append(item)
    return {
        "file": str(path),
        "slide_count": len(prs.slides),
        "width_in": round(prs.slide_width / EMU, 3),
        "height_in": round(prs.slide_height / EMU, 3),
        "slides": slides,
    }


def markdown(data):
    lines = [
        "# Template inventory",
        "",
        "- File: `%s`" % data["file"],
        "- Slides: %d" % data["slide_count"],
        "- Canvas: %s x %s in" % (data["width_in"], data["height_in"]),
        "",
        "| # | Layout | Archetype | Title | Shapes | Pics | Charts | Tables | Arrows | Notes |",
        "|---:|---|---|---|---:|---:|---:|---:|---:|:---:|",
    ]
    for s in data["slides"]:
        title = s["title"].replace("|", "/").replace("\n", " ")
        row = dict(s)
        row["title"] = title
        row["notes"] = "yes" if s["has_notes"] else ""
        lines.append("| {index} | {layout} | {archetype} | {title} | {shape_count} | {picture_count} | {chart_count} | {table_count} | {arrow_count} | {notes} |".format(**row))
    lines.append("")
    return "\n".join(lines)


def main():
    ap = argparse.ArgumentParser(description="Inventory a PPTX template before authoring.")
    ap.add_argument("pptx")
    ap.add_argument("--json")
    ap.add_argument("--markdown")
    args = ap.parse_args()
    data = inventory(args.pptx)
    if args.json:
        Path(args.json).write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")
    md = markdown(data)
    if args.markdown:
        Path(args.markdown).write_text(md, encoding="utf-8")
    print(md)


if __name__ == "__main__":
    main()
