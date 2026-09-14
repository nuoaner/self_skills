#!/usr/bin/env python3
import argparse
import re
import sys
from collections import Counter
from pptx import Presentation

DEFAULT_PATTERNS = [
    r"\blorem\b",
    r"\bipsum\b",
    r"\bTODO\b",
    r"\bTBD\b",
    r"\[\s*insert[^\]]*\]",
    r"\bxxx+\b",
    r"placeholder",
]


def slide_text(slide):
    parts = []
    for shape in slide.shapes:
        if getattr(shape, "has_text_frame", False):
            parts.append(shape.text or "")
        if getattr(shape, "has_table", False):
            for row in shape.table.rows:
                for cell in row.cells:
                    parts.append(cell.text or "")
    return "\n".join(parts).strip()


def slide_title(slide):
    if slide.shapes.title is not None:
        text = (slide.shapes.title.text or "").strip()
        if text:
            return text
    candidates = []
    for shape in slide.shapes:
        if not getattr(shape, "has_text_frame", False):
            continue
        text = (shape.text or "").strip()
        if not text:
            continue
        candidates.append((shape.top, text))
    return min(candidates)[1] if candidates else ""


def notes_text(slide):
    try:
        tf = slide.notes_slide.notes_text_frame
        return (tf.text or "").strip() if tf is not None else ""
    except Exception:
        return ""


def main():
    ap = argparse.ArgumentParser(description="Audit PPTX visible text for common leftovers and content hygiene issues.")
    ap.add_argument("pptx")
    ap.add_argument("--allow", action="append", default=[])
    ap.add_argument("--strict-notes", action="store_true")
    args = ap.parse_args()

    prs = Presentation(args.pptx)
    findings = []
    titles = []
    allow = set(args.allow)
    compiled = [(p, re.compile(p, re.IGNORECASE)) for p in DEFAULT_PATTERNS]

    for idx, slide in enumerate(prs.slides, 1):
        text = slide_text(slide)
        title = slide_title(slide)
        titles.append(title)
        if not text and len(slide.shapes) == 0:
            findings.append(("ERROR", idx, "empty slide"))
        elif not text:
            findings.append(("WARN", idx, "slide has no visible text"))
        for raw, pat in compiled:
            if raw in allow:
                continue
            if pat.search(text):
                findings.append(("ERROR", idx, "placeholder pattern matched: %s" % raw))
        if args.strict_notes and idx > 1 and not notes_text(slide):
            findings.append(("WARN", idx, "speaker notes are empty"))

    title_counts = Counter(t for t in titles if t)
    for title, count in title_counts.items():
        if count > 2:
            findings.append(("WARN", 0, "title repeated %d times: %s" % (count, title)))

    print("Slides: %d" % len(prs.slides))
    print("Findings: %d" % len(findings))
    for level, idx, message in findings:
        prefix = "Slide %d" % idx if idx else "Deck"
        print("%s %s: %s" % (level, prefix, message))

    if any(level == "ERROR" for level, _, _ in findings):
        sys.exit(1)


if __name__ == "__main__":
    main()
